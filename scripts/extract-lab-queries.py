#!/usr/bin/env python3
"""Витяг SQL зі студентських звітів для перевірки міри схожості (З7).

На вхід — теки `labs/labs/<щось>_<назва лаби>/` з роботами студентів (docx, pdf,
sql), вивантажені з LMS. На вихід — `labs/labs/_extracted/queries.jsonl`:
по одному рядку на запит, `{lab, lab_dir, student, source, sql}`, плюс дамп
`labs/labs/_extracted/<lab>.md` для перегляду очима і зведення покриття.

Скрипт нічого не тлумачить — лише дістає текст запитів. Рішення, чи міра
відрізняє лабу від лаби, ухвалює `app.shapecheck_files` і зрештою людина.

Межа методу: SQL, вставлений у звіт КАРТИНКОЮ (скріншот), сюди не потрапляє —
OCR тут немає навмисно. Скільки студентів через це випало — у зведенні.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "front"))
from app.sqlshape import shape  # noqa: E402

# Тека з роботами студентів. Перший аргумент або labs/labs.
SRC_DIR = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "labs" / "labs"
OUT_DIR = SRC_DIR / "_extracted"

# Перше слово стейтмента. `selct`/`slect` — часті опечатки, і саме такі запити
# викладачеві потрібні найбільше, тож ловимо і їх.
_START = re.compile(
    r"(?i)^\s*(with|select|sel[e]?ct|slect|insert|update|delete|create|alter|drop"
    r"|truncate|grant|revoke|merge|copy|explain|call|comment\s+on)\b"
)
# Рядок, який явно є прозою (українське речення), а не продовженням SQL.
_PROSE = re.compile(r"[а-яґєіїА-ЯҐЄІЇ]{4}.*[а-яґєіїА-ЯҐЄІЇ]{4}")
_KEYWORDS_LINE = re.compile(
    r"(?i)\b(from|where|join|group\s+by|order\s+by|having|select|values|set|into"
    r"|on|and|or|union|limit|returning|as|when|then|else|end|primary\s+key"
    r"|foreign\s+key|references|constraint|default|not\s+null)\b"
)
MAX_LINES = 60


def _statements(text: str) -> list[str]:
    """Вибирає стейтменти з довільного тексту: від ключового першого слова до
    `;`, порожнього рядка, наступного стейтмента або переходу в прозу.

    Глибина дужок рахується по кожному рядку: поки вона > 0, ми всередині
    незакритого підзапиту (`WHERE x > (` на своєму рядку, тіло підзапиту
    нижче), і межі стейтмента (порожній рядок, новий `SELECT`, проза) там не
    діють — інакше багаторядковий вкладений підзапит розривається навпіл:
    зовнішній запит обривається на відкритій `(`, а внутрішній `SELECT`
    підхоплюється як «новий» стейтмент і дожинає до чужої `;` з осиротілим
    хвостом `)`. Обидва фрагменти реальні в даних (докладніше — задача З19,
    аналіз позицій 3/7/9 у Lab 05 показав саме цю картину)."""
    lines = text.replace("\r", "").split("\n")
    out: list[str] = []
    i = 0
    while i < len(lines):
        if not _START.match(lines[i]):
            i += 1
            continue
        buf = [lines[i]]
        depth = lines[i].count("(") - lines[i].count(")")
        i += 1
        while i < len(lines) and len(buf) < MAX_LINES:
            ln = lines[i]
            if ";" in ln and depth <= 0:
                buf.append(ln[: ln.index(";") + 1])
                i += 1
                break
            if depth <= 0:
                if not ln.strip():
                    break
                if _START.match(ln):
                    break
                # Проза без жодного SQL-слова — кінець стейтмента.
                if _PROSE.search(ln) and not _KEYWORDS_LINE.search(ln):
                    break
            buf.append(ln)
            depth += ln.count("(") - ln.count(")")
            i += 1
        joined = re.sub(r"\s+", " ", " ".join(buf)).strip()
        # Кілька стейтментів в одному абзаці (`INSERT …; INSERT …`) — розбити.
        for k, part in enumerate(joined.split(";")):
            part = part.strip()
            if part and (k == 0 or _START.match(part)):
                out.append(part)
    return out


def _keep(sql: str) -> bool:
    if not (8 <= len(sql) <= 4000):
        return False
    sh = shape(sql)
    return bool(sh["tables"] or sh["constructs"] or sh["funcs"]
                or len(sh["columns"]) >= 2)


def _norm(sql: str) -> str:
    return re.sub(r"\s+", " ", sql).strip().lower()


def _from_docx(path: Path) -> str:
    import docx  # noqa: PLC0415

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _from_pdf(path: Path) -> str:
    import pdfplumber  # noqa: PLC0415

    with pdfplumber.open(str(path)) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _from_sql(path: Path) -> str:
    return path.read_text(errors="replace")


READERS = {".docx": _from_docx, ".pdf": _from_pdf, ".sql": _from_sql}

# id-префікс LMS: `1303325-Звіт.docx`. Прізвище з імені файлу — надійніший ключ
# студента, ніж id (docx і pdf тієї самої роботи мають різні id).
_ID = re.compile(r"^(\d{6,})-(.*)$")
_SURNAME = re.compile(r"[А-ЯҐЄІЇ][а-яґєії']{3,}")


def _student_key(name: str) -> str:
    m = _ID.match(name)
    ident, rest = (m.group(1), m.group(2)) if m else ("?", name)
    stem = Path(rest).stem
    surn = _SURNAME.findall(stem)
    bad = {"Лабораторна", "Робота", "Звіт", "Оновлений", "Бази", "Даних",
           "Лаб", "Письмовий", "Індекси", "Тригери", "Транзакції"}
    surn = [s for s in surn if s not in bad]
    return surn[0] if surn else ident


def main() -> int:
    if not SRC_DIR.is_dir():
        print(f"немає {SRC_DIR}", file=sys.stderr)
        return 2
    OUT_DIR.mkdir(exist_ok=True)

    rows: list[dict] = []
    summary: list[tuple] = []
    for lab_dir in sorted(p for p in SRC_DIR.iterdir() if p.is_dir() and p.name != "_extracted"):
        lab = lab_dir.name.split("_", 1)[-1].strip()
        students: dict[str, set[str]] = {}
        per_student_seen: dict[str, set[str]] = {}
        files = [f for f in sorted(lab_dir.iterdir()) if f.suffix.lower() in READERS]
        n_files_ok = 0
        for f in files:
            try:
                text = READERS[f.suffix.lower()](f)
            except Exception as e:  # noqa: BLE001
                print(f"  ! {f.name}: {e}", file=sys.stderr)
                continue
            n_files_ok += 1
            key = _student_key(f.name)
            seen = per_student_seen.setdefault(key, set())
            got = students.setdefault(key, set())
            for stmt in _statements(text):
                if not _keep(stmt):
                    continue
                nk = _norm(stmt)
                if nk in seen:
                    continue
                seen.add(nk)
                got.add(stmt)
                rows.append({"lab": lab, "lab_dir": lab_dir.name,
                             "student": key, "source": f.name, "sql": stmt})
        with_q = sum(1 for v in students.values() if v)
        summary.append((lab, len(files), n_files_ok, len(students), with_q,
                        sum(len(v) for v in students.values())))

        dump = OUT_DIR / f"{lab_dir.name}.md"
        with dump.open("w") as fh:
            fh.write(f"# {lab}\n\n")
            for key, qs in sorted(students.items()):
                fh.write(f"## {key}  ({len(qs)})\n\n")
                for q in sorted(qs):
                    fh.write(f"    {q}\n")
                fh.write("\n")

    out = OUT_DIR / "queries.jsonl"
    with out.open("w") as fh:
        for r in rows:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")

    print(f"{'лаба':42s} {'файлів':>7} {'ок':>4} {'студ':>5} {'з SQL':>6} {'запитів':>8}")
    for lab, nf, ok, ns, wq, nq in summary:
        flag = "  ⚠ мало" if wq < 3 else ""
        print(f"{lab[:42]:42s} {nf:7d} {ok:4d} {ns:5d} {wq:6d} {nq:8d}{flag}")
    tot_s = sum(s[3] for s in summary)
    tot_wq = sum(s[4] for s in summary)
    print(f"\n{len(rows)} запитів, {tot_wq} з {tot_s} студентів дали хоч один")
    print(f"студентів без жодного запиту (ймовірно скріншоти): {tot_s - tot_wq}")
    print(f"→ {out}")
    print(f"→ дампи на лабу: {OUT_DIR}/<тека>.md")
    return 0


if __name__ == "__main__":
    sys.exit(main())
